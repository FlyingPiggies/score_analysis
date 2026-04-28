from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .constants import (
    COLUMN_PADDING_WIDTH,
    FREEZE_PANES_ANCHOR,
    MAX_COLUMN_WIDTH,
    MIN_COLUMN_WIDTH,
    SHEET_NAME_COMPARISON,
    SHEET_NAME_TOP,
)
from .text_utils import clean_text


def write_workbook(output_path: Path, comparison_sheet_df: pd.DataFrame, top_sheet_df: pd.DataFrame) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        comparison_sheet_df.to_excel(writer, sheet_name=SHEET_NAME_COMPARISON, index=False)
        top_sheet_df.to_excel(writer, sheet_name=SHEET_NAME_TOP, index=False)
    autosize_columns(output_path)


def autosize_columns(workbook_path: Path) -> None:
    workbook = load_workbook(workbook_path)
    for sheet_name in workbook.sheetnames:
        worksheet = workbook[sheet_name]
        worksheet.freeze_panes = FREEZE_PANES_ANCHOR
        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = column_cells[0].column_letter
            for cell in column_cells:
                cell_text = clean_text(cell.value)
                if len(cell_text) > max_length:
                    max_length = len(cell_text)
            worksheet.column_dimensions[column_letter].width = max(
                MIN_COLUMN_WIDTH,
                min(max_length + COLUMN_PADDING_WIDTH, MAX_COLUMN_WIDTH),
            )
    workbook.save(workbook_path)
